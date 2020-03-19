import openpyxl

def write_data(sheet_, wb, event_session_id, link, room, row, path):
    sheet_.cell(row = row, column = 15).value = event_session_id
    sheet_.cell(row = row, column = 9).value = link
    sheet_.cell(row = row, column = 10).value = room
    wb.save(path)
